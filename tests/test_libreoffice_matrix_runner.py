from __future__ import annotations

import json
import stat
import subprocess
import sys
import tempfile
import textwrap
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
RUNNER = ROOT / "scripts" / "application_gates" / "run_libreoffice_matrix.py"
GENERATOR = ROOT / "scripts" / "application_gates" / "generate_fixtures.py"
FAKE = rf"""#!{sys.executable}
import pathlib, shutil, sys
out = pathlib.Path(sys.argv[sys.argv.index("--outdir") + 1])
source = pathlib.Path(sys.argv[-1])
out.mkdir(parents=True, exist_ok=True)
if source.suffix.lower() == ".xlsx":
    shutil.copyfile(source, out / (source.stem + ".xlsx"))
else:
    (out / (source.stem + ".pdf")).write_bytes(b"%PDF-1.4\n% compatibility fixture\n")
"""


class LibreOfficeMatrixRunnerTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.fixtures = self.root / "fixtures"
        self.fixtures.mkdir()
        Document().save(self.fixtures / "sample.docx")
        Workbook().save(self.fixtures / "sample.xlsx")
        Presentation().save(self.fixtures / "sample.pptx")
        self.executable = self.root / "soffice"
        self.executable.write_text(textwrap.dedent(FAKE))
        self.executable.chmod(self.executable.stat().st_mode | stat.S_IXUSR)
        self.manifest = self.fixtures / "manifest.json"
        self.manifest.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "fixtures": [
                        {"id": "docx-basic", "artifact_type": "docx", "path": "sample.docx"},
                        {"id": "xlsx-basic", "artifact_type": "xlsx", "path": "sample.xlsx"},
                        {"id": "pptx-basic", "artifact_type": "pptx", "path": "sample.pptx"},
                    ],
                }
            )
        )

    def tearDown(self) -> None:
        self.temp.cleanup()

    def run_matrix(self, manifest: Path | None = None) -> tuple[subprocess.CompletedProcess[str], dict]:
        completed = subprocess.run(
            [
                sys.executable,
                str(RUNNER),
                "--manifest",
                str(manifest or self.manifest),
                "--workdir",
                str(self.root / "work"),
                "--executable",
                str(self.executable),
                "--runtime-version",
                "LibreOffice Fake 1.0",
                "--runtime-image-digest",
                "sha256:" + "b" * 64,
            ],
            text=True,
            capture_output=True,
            check=False,
        )
        return completed, json.loads(completed.stdout)

    def test_matrix_records_exact_runtime_and_all_fixture_reports(self) -> None:
        completed, report = self.run_matrix()
        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertEqual(report["status"], "ok")
        self.assertEqual(report["runtime_identity"]["image_digest"], "sha256:" + "b" * 64)
        self.assertEqual([item["id"] for item in report["fixtures"]], ["docx-basic", "xlsx-basic", "pptx-basic"])
        self.assertTrue(all(item["report"]["status"] == "ok" for item in report["fixtures"]))
        self.assertFalse(report["microsoft_office_equivalence_claimed"])

    def test_manifest_path_escape_is_refused_before_execution(self) -> None:
        escaped = self.fixtures / "escaped.json"
        escaped.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "fixtures": [{"id": "escape", "artifact_type": "docx", "path": "../outside.docx"}],
                }
            )
        )
        completed, report = self.run_matrix(escaped)
        self.assertEqual(completed.returncode, 2)
        self.assertEqual(report, {"schema_version": 1, "status": "refused", "reason": "invalid_manifest"})

    def test_generator_creates_closed_three_format_fixture_manifest(self) -> None:
        output = self.root / "generated"
        completed = subprocess.run(
            [sys.executable, str(GENERATOR), "--output-dir", str(output)],
            text=True,
            capture_output=True,
            check=False,
        )
        self.assertEqual(completed.returncode, 0, completed.stderr)
        manifest = json.loads((output / "manifest.json").read_text())
        self.assertEqual(manifest["schema_version"], 1)
        self.assertEqual(
            [(item["artifact_type"], item["path"]) for item in manifest["fixtures"]],
            [("docx", "basic.docx"), ("xlsx", "recalculation.xlsx"), ("pptx", "basic.pptx")],
        )
        self.assertTrue(all((output / item["path"]).is_file() for item in manifest["fixtures"]))
        cases = json.loads((output / "recalculation-cases.json").read_text())
        self.assertEqual(cases["cells"][0]["cell"], "B1")
        self.assertEqual(cases["cells"][0]["expected_value"], 2)

    def test_real_libreoffice_workflow_is_manual_and_requires_pinned_host(self) -> None:
        workflow = (ROOT / ".github" / "workflows" / "libreoffice-compat.yml").read_text()
        self.assertIn("workflow_dispatch:", workflow)
        self.assertNotIn("schedule:", workflow)
        self.assertIn("office-libreoffice-pinned", workflow)
        self.assertIn("LIBREOFFICE_IMAGE_DIGEST", workflow)
        self.assertIn("run_libreoffice_matrix.py", workflow)
        self.assertIn("run_xlsx_recalculation.py", workflow)
        self.assertNotIn("rm -rf", workflow)
        self.assertIn("github.run_attempt", workflow)


if __name__ == "__main__":
    unittest.main()
